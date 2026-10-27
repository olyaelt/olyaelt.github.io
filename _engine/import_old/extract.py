"""Dump slides with hidden (highlighted) text marked as [[...]]; usage: extract.py deck.pptx 15 16 ..."""
import sys
from pptx import Presentation
from lxml import etree
NS = {'a':'http://schemas.openxmlformats.org/drawingml/2006/main'}
def run_hidden(r):
    rpr = r.find('a:rPr', NS)
    if rpr is None: return None
    hl = rpr.find('a:highlight', NS)
    if hl is not None:
        c = hl.find('.//a:srgbClr', NS)
        return 'H' + (c.get('val') if c is not None else '?')
    return None
def color(r):
    rpr = r.find('a:rPr', NS)
    if rpr is None: return ''
    c = rpr.find('a:solidFill/a:srgbClr', NS)
    return c.get('val') if c is not None else ''
prs = Presentation(sys.argv[1])
for n in map(int, sys.argv[2:]):
    s = prs.slides[n-1]
    print(f'################ slide {n}')
    for sh in s.shapes:
        if not sh.has_text_frame: continue
        out=[]
        for p in sh.text_frame._txBody.findall('a:p', NS):
            line=''
            for r in p:
                tag = etree.QName(r).localname
                if tag == 'r':
                    t = r.findtext('a:t', namespaces=NS) or ''
                    h = run_hidden(r)
                    col = color(r)
                    if h: line += f'[[{t}]]' if h in ('HFFFF00',) else f'[[{h}:{t}]]'
                    elif col in ('0000FF','0000ff'): line += f'{{{t}}}'
                    else: line += t
                elif tag == 'br': line += ' / '
            out.append(line)
        txt = '\n'.join(out).strip()
        if txt: print('--', sh.name, '\n' + txt)
