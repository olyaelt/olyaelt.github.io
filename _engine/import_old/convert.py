"""Turn the homework slides (yellow-on-yellow answers, blue first letters, orange IPA) into page data.
   python3 convert.py  -> writes gaps16.json, gaps21.json"""
import json, re
from pptx import Presentation
from lxml import etree
NS = {'a': 'http://schemas.openxmlformats.org/drawingml/2006/main'}
CYR = re.compile(r'[А-Яа-яЁё]')

def runs_of_par(p):
    """[(kind, text)] kind: plain | blue | hid | ipa"""
    out = []
    for r in p:
        tag = etree.QName(r).localname
        if tag == 'br': out.append(('plain', ' ')); continue
        if tag != 'r': continue
        t = r.findtext('a:t', namespaces=NS) or ''
        rpr = r.find('a:rPr', NS)
        kind = 'plain'
        if rpr is not None:
            hl = rpr.find('a:highlight', NS)
            col = rpr.find('a:solidFill/a:srgbClr', NS)
            if hl is not None:
                c = hl.find('.//a:srgbClr', NS)
                kind = 'ipa' if (c is not None and c.get('val').upper() == 'FF9900') else 'hid'
            elif col is not None and col.get('val').upper() == '0000FF':
                kind = 'blue'
        if out and out[-1][0] == kind and kind in ('plain', 'blue', 'hid'):
            out[-1] = (kind, out[-1][1] + t)
        else:
            out.append((kind, t))
    return out

def parts_of(runs):
    """runs -> engine parts: strings and {f, r} gaps ({f:'', r, o:1} for hidden IPA)"""
    parts = []
    def add_text(s):
        if not s: return
        if parts and isinstance(parts[-1], str): parts[-1] += s
        else: parts.append(s)
    i = 0
    while i < len(runs):
        kind, t = runs[i]
        nxt = runs[i + 1] if i + 1 < len(runs) else None
        if kind == 'plain':
            if nxt and nxt[0] == 'hid':
                m = re.search(r"([A-Za-z'’]+)$", t)
                if m:
                    add_text(t[:m.start()]); parts.append({'f': m.group(1), 'r': nxt[1].strip()}); i += 2; continue
            add_text(t)
        elif kind == 'blue':
            if nxt and nxt[0] == 'hid':
                parts.append({'f': t.strip(), 'r': nxt[1].strip()}); i += 2; continue
            if re.search(r'[A-Za-z]', t): parts.append({'f': t.strip(), 'r': ''})
            else: add_text(t)
        elif kind == 'hid':
            parts.append({'f': '', 'r': t.strip()})
        elif kind == 'ipa':
            parts.append({'f': '', 'r': t.strip(), 'o': 1})
        i += 1
    # spaces between neighbouring gaps are separate strings, as in the engine's data
    out = []
    for x in parts:
        if isinstance(x, str):
            x = re.sub(r'[  ]{2,}', ' ', x)
            if x: out.append(x)
        else: out.append(x)
    return out

def strip_ends(parts):
    if parts and isinstance(parts[0], str): parts[0] = parts[0].lstrip()
    if parts and isinstance(parts[-1], str): parts[-1] = parts[-1].rstrip()
    return [p for p in parts if p != '']

def text_of(parts): return ''.join(p if isinstance(p, str) else (p['f'] + p['r']) for p in parts)

def shape_pars(slide, pick):
    """paragraph run-lists of the shape chosen by pick(text)"""
    for sh in slide.shapes:
        if not sh.has_text_frame: continue
        txt = sh.text_frame.text
        if pick(txt):
            return [runs_of_par(p) for p in sh.text_frame._txBody.findall('a:p', NS)]
    raise SystemExit('shape not found')

def dialogue(slide, first_speakers):
    """(intro, lines) from the text shape that has the speakers"""
    pars = shape_pars(slide, lambda t: any(re.search(r'(^|\n)' + s + r':', t) for s in first_speakers))
    intro, lines = '', []
    for runs in pars:
        parts = strip_ends(parts_of(runs))
        if not parts: continue
        first = parts[0] if isinstance(parts[0], str) else ''
        m = re.match(r'^([A-Z][a-z]+): ?', first)
        if m:
            parts[0] = first[m.end():]
            lines.append({'s': m.group(1), 'p': strip_ends(parts)})
        elif not lines:
            intro = text_of(parts).strip()
        else:
            raise SystemExit('stray paragraph: ' + text_of(parts))
    return intro, lines

def glossary(slide, pick):
    entries = []
    for runs in shape_pars(slide, pick):
        parts = strip_ends(parts_of(runs))
        if not parts: continue
        plain = text_of(parts)
        if CYR.search(plain) and not re.search(r'[A-Za-z]', plain):       # a Russian line of its own: belongs to the entry above
            entries[-1]['ru'] = plain.strip(); continue
        e = {'p': [], 'm': '', 'ru': ''}
        for x in parts:
            if isinstance(x, str) and (' — ' in x or e['m']):
                if not e['m']:
                    a, b = x.split(' — ', 1); e['p'].append(a); rest = b
                else:
                    rest = x
                k = CYR.search(rest)
                if k: e['m'] += rest[:k.start()]; e['ru'] += rest[k.start():]
                else: e['m'] += rest
            elif isinstance(x, str) and CYR.search(x) and not e['m']:
                k = CYR.search(x); e['p'].append(x[:k.start()]); e['ru'] += x[k.start():]
            else:
                (e['p'] if not e['ru'] else e['p']).append(x)
        e['p'] = strip_ends(e['p']); e['m'] = re.sub(r'\s+', ' ', e['m']).strip(); e['ru'] = re.sub(r'\s+', ' ', e['ru']).strip()
        entries.append(e)
    return entries

def phrase_box(slide):
    items = []
    for runs in shape_pars(slide, lambda t: 'ONE PHRASE IS EXTRA' in t):
        t = ''.join(x[1] for x in runs).strip()
        if not t or t == 'ONE PHRASE IS EXTRA': continue
        en, ru = [s.strip() for s in re.split(r'\s+[–-]\s+', t, 1)]
        items.append([en, ru])
    return items

L12 = Presentation('L12.pptx'); L34 = Presentation('L34.pptx')
S = lambda prs, n: prs.slides[n - 1]
d16 = {}
# 14.10: the two texts (glossary slide, then without)
for key, gl, plain, spk in (('coffee', 15, 16, ['Paul']), ('later', 17, 18, ['Paul'])):
    intro, lines = dialogue(S(L12, plain), spk)
    intro_g, lines_g = dialogue(S(L12, gl), spk)
    assert json.dumps(lines) == json.dumps(lines_g), key + ': glossary and plain versions differ'
    d16[key] = {'intro': intro, 'lines': lines, 'gloss': glossary(S(L12, gl), lambda t: ' — ' in t and 'Paul:' not in t)}
# 16.10: four conversations, with the box (one or two slides) and with first letters
for key, box_slides, fl, spk in (('conv1', [23, 24], 25, ['Lena', 'Mark']), ('conv2', [26], 27, ['Sofia']),
                                 ('conv3', [28, 29], 30, ['Paula', 'Igor']), ('conv4', [31, 32], 33, ['Tom', 'Vika'])):
    intro, lines_box = '', []
    for n in box_slides:
        i, l = dialogue(S(L12, n), spk); intro = intro or i; lines_box += l
    intro_f, lines_f = dialogue(S(L12, fl), spk)
    assert [x['s'] for x in lines_box] == [x['s'] for x in lines_f], key + ': speakers differ'
    d16[key] = {'intro': intro_f, 'box_lines': lines_box, 'lines': lines_f, 'box': phrase_box(S(L12, box_slides[0]))}
d21 = {}
for key, gl, plain, spk in (('dan', 35, 36, ['Sophie']), ('ruth', 38, 39, ['Sophie'])):
    intro, lines = dialogue(S(L34, plain), spk)
    _, lines_g = dialogue(S(L34, gl), spk)
    d21[key] = {'intro': intro, 'lines': lines, 'lines_g': lines_g,
                'gloss': glossary(S(L34, gl), lambda t: ('Sophie:' not in t) and ('—' in t))}
json.dump(d16, open('gaps16.json', 'w'), ensure_ascii=False, indent=1)
json.dump(d21, open('gaps21.json', 'w'), ensure_ascii=False, indent=1)
for k, v in list(d16.items()) + list(d21.items()):
    print('==', k, '|', v['intro'][:70])
    for l in v.get('box_lines', [])[:2]: print('  box', l['s'], l['p'])
    for l in v['lines'][:2]: print('  ', l['s'], l['p'])
    for g in v.get('gloss', [])[:3]: print('  gl', g)
    if 'box' in v: print('  box', v['box'])
