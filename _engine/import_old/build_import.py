"""The C1 import file: old spreadsheet ticks → page tasks, the student's texts with Olya's corrections.
   python3 build_import.py  -> 'Import old ticks C1.xlsx' + a printout to check"""
import json, re
import openpyxl
from openpyxl.cell.rich_text import CellRichText, TextBlock

SRC = '/mnt/user-data/attachments/1bcc0e56-36de-58a4-bfbf-17733b9d9d19/0e9c83eb-GE_Wed_Fri_20.15_2026-27_1.xlsx'
FR = json.load(open('frames.json'))
P16, P21, P23 = 'hw-ge-c1-2026-10-16', 'hw-ge-c1-2026-10-21', 'hw-ge-c1-2026-10-23'
# spreadsheet row → (page, [tasks], in the old count?)
ROWS = {
    61: (P16, ['coffee-text', 'later-text'], True),   # 14.10 revise the expressions on slides 15–18
    42: (P16, ['practice'], True),                    # 16.10 slides 41–70
    44: (P16, ['convos'], True),                      # 16.10 slides 22–34
    47: (P16, ['listen1'], True), 48: (P16, ['listen2'], True), 49: (P16, ['listen3'], True),
    51: (P16, ['write'], True), 52: (P16, ['write'], True),
    31: (P21, ['traps'], True),                       # 21.10 slide 31
    32: (P21, ['dan-text', 'ruth-text'], True),       # 21.10 slides 34–39
    34: (P21, ['yg-get'], True), 35: (P21, ['yg-easy'], True), 36: (P21, ['yg-training'], True),
    6: (P23, ['clocks-text'], False),                 # 23.10 slide 46 or Brainscape
    10: (P23, ['clock-listen'], False),               # 23.10 TubeQuizard quiz 1
    8: (P23, ['career-listen'], False),
    12: (P23, ['career-rewrite'], False),             # 23.10 slides 122–143 (the 8 career sentences; the 2 clock ones sit in a mixed task)
    13: (P23, ['career-finish'], False),
    17: (P23, ['write'], False),
}
NOT_DONE = re.compile(r'not yet|not available', re.I)

wb = openpyxl.load_workbook(SRC, rich_text=True); ws = wb['Homework']
# names as they are on the Students tab of her table (the old spreadsheet used other short forms)
NAME_MAP = {'Julia': 'Julia Ch', 'Petr S': 'Petr Sk', 'Peter Smirnov': 'Peter Sm', 'Artyom': 'Tioma'}
names = {c.column: NAME_MAP.get(c.value, c.value) for c in ws[1][3:12]}
def done(cell):
    v = cell.value
    fill = cell.fill.fgColor.rgb if cell.fill is not None and cell.fill.fill_type == 'solid' else None
    if v not in (None, '') and NOT_DONE.search(str(v)): return False
    return bool(v not in (None, '')) or (fill not in (None, '00000000', 'FFFFFFFF'))
ticks = {}
for r, (page, tasks, old) in ROWS.items():
    for col, name in names.items():
        if done(ws.cell(r, col)):
            for t in tasks: ticks[(name, page, t)] = old

# the student's texts in the sheet: bold = good, strikethrough = to change, yellow = Olya's words in [brackets]
def runs(cell):
    v = cell.value; out = []
    for b in (v if isinstance(v, CellRichText) else [v]):
        if isinstance(b, TextBlock):
            f = b.font; col = f.color.rgb if f is not None and f.color is not None and f.color.type == 'rgb' else None
            kind = 'comment' if col == 'FFFFFF00' else 'strike' if (f is not None and f.strike) else 'bold' if (f is not None and f.b) else 'plain'
            out.append([kind, b.text])
        else: out.append(['plain', str(b)])
    merged = []
    for k, t in out:
        if merged and merged[-1][0] == k: merged[-1][1] += t
        else: merged.append([k, t])
    for i, x in enumerate(merged):          # brackets around Olya's words belong to her words
        k, t = x[0], x[1]
        if k != 'comment': continue
        lead = ' ' if re.match(r'^\s', t) else ''
        merged[i][1] = t.strip().strip('[]').strip()
        merged[i].append(lead)
        if i > 0: merged[i - 1][1] = re.sub(r'\s*\[\s*$', ' ', merged[i - 1][1])
        if i + 1 < len(merged): merged[i + 1][1] = re.sub(r'^\s*\]', '', merged[i + 1][1])
    return merged
def wrap(t, a, b):          # mark each line separately; keep the spaces outside the marks
    out = []
    for line in t.split('\n'):
        m = re.match(r'^(\s*(?:-\s+)?)(.*?)(\s*)$', line)
        out.append(m.group(1) + (a + m.group(2) + b if m.group(2) else '') + m.group(3))
    return '\n'.join(out)
def tidy(s): return re.sub(r'[ \t]+([.,?!])', r'\1', re.sub(r'[ \t]{2,}', ' ', s)).strip()
def student_text(rs): return tidy(''.join(x[1] for x in rs if x[0] != 'comment'))
def feedback(rs):
    s = ''.join(wrap(x[1], '**', '**') if x[0] == 'bold' else wrap(x[1], '~~', '~~') if x[0] == 'strike' else (x[2] + '{' + x[1] + '}' if x[0] == 'comment' else x[1]) for x in rs)
    s = s.replace('** **', ' ')
    return re.sub(r'[ \t]{2,}', ' ', s).strip()
r51, r52 = runs(ws['G51']), runs(ws['G52'])
TEXT = '1) ' + student_text(r51) + '\n\n2) ' + student_text(r52)
FEEDBACK = '1) ' + feedback(r51) + '\n\n2) ' + feedback(r52)

title = {k: {t['id']: t for t in v['tasks']} for k, v in FR.items()}
out = openpyxl.Workbook(); sh = out.active; sh.title = 'Import old ticks C1'
sh.append(['Type', 'Student', 'Group', 'Homework', 'Date', 'Task', 'Title', 'Frames', 'In old count', 'Text', 'Feedback'])
for page in (P23, P21, P16):
    for t in FR[page]['tasks']: sh.append(['task', '', 'C1', page, FR[page]['date'], t['id'], t['title'], t['frames'], '', '', ''])
order = {P16: 0, P21: 1, P23: 2}
old_frames = 0
for (name, page, task), old in sorted(ticks.items(), key=lambda x: (order[x[0][1]], x[0][2], x[0][0])):
    t = title[page][task]
    sh.append(['tick', name, 'C1', page, FR[page]['date'], task, t['title'], t['frames'], 'yes' if old else 'no', '', ''])
    old_frames += t['frames'] if old else 0
sh.append(['text', names[7], 'C1', P16, FR[P16]['date'], 'write', title[P16]['write']['title'], '', '', TEXT, FEEDBACK])
for col, w in zip('ABCDEFGHIJK', (6, 14, 6, 22, 9, 14, 34, 7, 10, 60, 60)): sh.column_dimensions[col].width = w
out.save('Import old ticks C1.xlsx')

print('ticks:', len(ticks), '| frames in the old count:', old_frames)
by = {}
for (name, page, task) in ticks: by.setdefault(name, []).append(page[-5:] + ' ' + task)
for n in sorted(by): print(f'  {n}: ' + ', '.join(sorted(by[n])))
print('\nTEXT:\n' + TEXT + '\n\nFEEDBACK:\n' + FEEDBACK)
